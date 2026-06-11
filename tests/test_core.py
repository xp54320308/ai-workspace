"""
第一层测试：单元测试
使用合成数据，隔离测试每个函数，无需外部依赖（Ollama/ChromaDB）。
"""

import os
import sys
import tempfile
import shutil
import pytest

# 确保项目根目录在 path 中
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


# ============================================================
#  ingest_docs 模块测试
# ============================================================

class TestDocumentSplitting:
    """测试文档切片逻辑"""

    def test_split_single_document(self):
        """单个文档切片"""
        from ingest_docs import split_documents
        from langchain_core.documents import Document

        doc = Document(
            page_content="第一段内容。\n\n第二段内容。\n\n" * 50,
            metadata={"source": "test.txt"},
        )
        chunks = split_documents([doc], chunk_size=200, chunk_overlap=50)

        assert len(chunks) > 1, "应该切出多个块"
        for chunk in chunks:
            assert len(chunk.page_content) <= 250, f"块太大: {len(chunk.page_content)}"
            assert chunk.metadata["source_file"] == "test.txt"

    def test_split_adds_source_file(self):
        """切片应自动添加 source_file 元数据"""
        from ingest_docs import split_documents
        from langchain_core.documents import Document

        doc = Document(
            page_content="测试内容 " * 100,
            metadata={"source": "/path/to/my_doc.txt"},
        )
        chunks = split_documents([doc])
        assert all(c.metadata["source_file"] == "my_doc.txt" for c in chunks)

    def test_split_empty_document(self):
        """空文档应返回空列表"""
        from ingest_docs import split_documents
        from langchain_core.documents import Document

        doc = Document(page_content="", metadata={"source": "empty.txt"})
        chunks = split_documents([doc])
        # 空内容可能产生 0 或 1 个块
        assert len(chunks) <= 1

    def test_split_preserves_content(self):
        """切片不应丢失内容"""
        from ingest_docs import split_documents
        from langchain_core.documents import Document

        text = "这是重要数据：合同编号 ABC-12345。" * 10
        doc = Document(page_content=text, metadata={"source": "contract.txt"})
        chunks = split_documents([doc], chunk_size=100, chunk_overlap=20)

        # 合并所有块的内容
        combined = "".join(c.page_content for c in chunks)
        assert "ABC-12345" in combined, "切片不应丢失关键数据"


class TestLoadDocx:
    """测试 .docx 加载"""

    def test_load_docx_creates_document(self):
        """load_docx 应返回 Document 对象列表"""
        from ingest_docs import load_docx

        # 创建临时 .docx 文件
        try:
            from docx import Document as DocxDocument
        except ImportError:
            pytest.skip("python-docx 未安装")

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            tmp_path = f.name

        try:
            doc = DocxDocument()
            doc.add_paragraph("测试段落一")
            doc.add_paragraph("测试段落二")
            doc.save(tmp_path)

            result = load_docx(tmp_path)
            assert len(result) == 1
            assert "测试段落一" in result[0].page_content
            assert "测试段落二" in result[0].page_content
            assert result[0].metadata["source"] == tmp_path
        finally:
            os.unlink(tmp_path)


class TestLoadXlsx:
    """测试 .xlsx 加载"""

    def test_load_xlsx_multiple_sheets(self):
        """多 Sheet 的 xlsx 应为每个 Sheet 创建一个 Document"""
        from ingest_docs import load_xlsx

        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl 未安装")

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp_path = f.name

        try:
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "Sheet1"
            ws1.append(["Name", "Value"])
            ws1.append(["ItemA", 100])

            ws2 = wb.create_sheet("Sheet2")
            ws2.append(["Product", "Price"])
            ws2.append(["Widget", 9.99])

            wb.save(tmp_path)

            result = load_xlsx(tmp_path)
            assert len(result) == 2
            sheet_names = [r.metadata["sheet"] for r in result]
            assert "Sheet1" in sheet_names
            assert "Sheet2" in sheet_names
        finally:
            os.unlink(tmp_path)

    def test_load_xlsx_skips_empty_rows(self):
        """空行应被跳过"""
        from ingest_docs import load_xlsx

        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl 未安装")

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            tmp_path = f.name

        try:
            wb = Workbook()
            ws = wb.active
            ws.append(["Header1", "Header2"])
            ws.append([None, None])  # 空行
            ws.append([None, None])  # 空行
            ws.append(["Data1", "Data2"])
            wb.save(tmp_path)

            result = load_xlsx(tmp_path)
            content = result[0].page_content
            # 空行不应出现在输出中
            lines = [l for l in content.split("\n") if l.strip() and not l.startswith("【")]
            assert len(lines) == 2  # Header + Data，两个空行被跳过
        finally:
            os.unlink(tmp_path)


# ============================================================
#  config 模块测试
# ============================================================

class TestConfig:
    """测试配置加载"""

    def test_default_config_loads(self):
        """默认配置应能正常加载"""
        from config import load_config

        config = load_config(config_path="/nonexistent/path.yaml")
        assert config["models"]["embedding"] == "nomic-embed-text"
        assert config["retrieval"]["k"] == 5
        assert config["retrieval"]["chunk_size"] == 800

    def test_env_override(self):
        """环境变量应覆盖默认值"""
        from config import load_config

        os.environ["RAG_CHAT_MODEL"] = "test-model-123"
        try:
            config = load_config(config_path="/nonexistent/path.yaml")
            assert config["models"]["chat"] == "test-model-123"
        finally:
            del os.environ["RAG_CHAT_MODEL"]

    def test_legacy_config_format(self):
        """旧格式 CONFIG 应正确转换"""
        from config import load_config, get_legacy_config

        config = load_config(config_path="/nonexistent/path.yaml")
        legacy = get_legacy_config(config)
        assert "persist_dir" in legacy
        assert "embedding_model" in legacy
        assert "chat_model" in legacy
        assert legacy["embedding_model"] == "nomic-embed-text"

    def test_get_config_value_dotpath(self):
        """点路径取值"""
        from config import load_config, get_config_value

        config = load_config(config_path="/nonexistent/path.yaml")
        # 使用不受环境变量影响的字段测试
        assert get_config_value(config, "retrieval.k") == 5
        assert get_config_value(config, "retrieval.chunk_size") == 800
        assert get_config_value(config, "nonexistent.path", "default") == "default"


# ============================================================
#  skill_manager 模块测试
# ============================================================

class TestSkillManager:
    """测试闭环学习 Skill 管理"""

    @pytest.fixture
    def skill_mgr(self, tmp_path):
        """创建临时 SkillManager"""
        from skill_manager import SkillManager
        return SkillManager(skills_dir=str(tmp_path / "skills"))

    def test_save_and_search(self, skill_mgr):
        """保存后应能搜索到"""
        skill_mgr.save_skill(
            question="什么是 RAG 技术？",
            answer="RAG 是检索增强生成技术...",
            sources=["knowledge.txt"],
            rating=5,
        )

        results = skill_mgr.search_skills("RAG 技术原理")
        assert len(results) > 0
        assert "RAG" in results[0]["question"]

    def test_search_no_match(self, skill_mgr):
        """不相关的查询应返回空"""
        skill_mgr.save_skill(
            question="Python 语法基础",
            answer="Python 使用缩进...",
        )
        results = skill_mgr.search_skills("量子物理前沿")
        assert len(results) == 0

    def test_get_skill_answer(self, skill_mgr):
        """读取 Skill 应返回完整内容"""
        skill_mgr.save_skill(
            question="ChromaDB 怎么用？",
            answer="ChromaDB 是一个向量数据库...",
            sources=["chromadb_guide.txt"],
        )

        # 获取 skill_id
        skills = skill_mgr.list_skills()
        assert len(skills) == 1
        skill_id = skills[0]["skill_id"]

        result = skill_mgr.get_skill_answer(skill_id)
        assert result is not None
        assert "向量数据库" in result["answer"]
        assert "chromadb_guide.txt" in result["sources"]

    def test_delete_skill(self, skill_mgr):
        """删除后应不可搜索"""
        skill_mgr.save_skill(question="临时问题", answer="临时答案")
        skills = skill_mgr.list_skills()
        skill_id = skills[0]["skill_id"]

        assert skill_mgr.delete_skill(skill_id) is True
        assert len(skill_mgr.list_skills()) == 0

    def test_stats(self, skill_mgr):
        """统计信息应正确"""
        skill_mgr.save_skill(question="Q1", answer="A1", rating=5)
        skill_mgr.save_skill(question="Q2", answer="A2", rating=3)

        stats = skill_mgr.get_stats()
        assert stats["total_skills"] == 2
        assert stats["avg_rating"] == 4.0

    def test_keywords_extraction(self, skill_mgr):
        """中英文关键词提取"""
        skill_mgr.save_skill(
            question="Python 的 LangChain 框架如何使用？",
            answer="LangChain 是一个强大的框架...",
        )

        results = skill_mgr.search_skills("LangChain framework")
        assert len(results) > 0


if __name__ == "__main__":
    pytest.main([__file__, "-v"])

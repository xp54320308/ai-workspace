"""
文档摄入管道 v2 - 支持元数据增强

新增功能：
1. 文档类型分类
2. 项目名称提取
3. 章节标题提取
4. 关键词提取
"""

import os
from typing import List, Optional
from langchain_community.document_loaders import (
    DirectoryLoader,
    TextLoader,
    PyPDFLoader,
)
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

from metadata_enricher import MetadataEnricher


# ---------------------------------------------------------------------------
# 文件格式加载器映射
# ---------------------------------------------------------------------------
LOADER_MAPPING = {
    ".txt":  {"loader_cls": TextLoader,   "loader_kwargs": {"encoding": "utf-8"}},
    ".pdf":  {"loader_cls": PyPDFLoader,  "loader_kwargs": {}},
    ".md":   {"loader_cls": TextLoader,   "loader_kwargs": {"encoding": "utf-8"}},
}


def load_docx(file_path: str) -> List[Document]:
    """使用 python-docx 加载 .docx 文件"""
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    return [Document(page_content=text, metadata={"source": file_path})]


def load_xlsx(file_path: str) -> List[Document]:
    """使用 openpyxl 加载 .xlsx 文件，所有Sheet转为文本"""
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    docs = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if all(cell is None for cell in row):
                continue
            line = "\t".join(str(cell) if cell is not None else "" for cell in row)
            rows.append(line)
        if rows:
            text = f"【工作表: {sheet_name}】\n" + "\n".join(rows)
            docs.append(Document(page_content=text, metadata={"source": file_path, "sheet": sheet_name}))
    wb.close()
    return docs


def load_single_document(file_path: str) -> List[Document]:
    """
    加载单个文件
    
    Args:
        file_path: 文件路径
    
    Returns:
        Document 列表
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext in LOADER_MAPPING:
        config = LOADER_MAPPING[ext]
        loader = config["loader_cls"](file_path, **config["loader_kwargs"])
        return loader.load()
    elif ext == ".docx":
        return load_docx(file_path)
    elif ext == ".xlsx":
        return load_xlsx(file_path)
    else:
        print(f"⚠️ 不支持的文件格式: {ext}")
        return []


def load_documents(data_dir: str = "data", enrich_metadata: bool = True) -> List[Document]:
    """
    从 data_dir 加载所有支持格式的文档
    
    Args:
        data_dir: 数据目录
        enrich_metadata: 是否增强元数据
    
    Returns:
        Document 列表
    """
    all_docs = []

    if not os.path.exists(data_dir):
        print(f"❌ Error: {data_dir} directory not found.")
        return []

    print(f"📂 Scanning directory: {os.path.abspath(data_dir)}")

    # 初始化元数据增强器
    enricher = MetadataEnricher() if enrich_metadata else None

    # 1. 加载 .txt, .pdf, .md 文件
    for ext, config in LOADER_MAPPING.items():
        try:
            loader = DirectoryLoader(
                data_dir,
                glob=f"**/*{ext}",
                loader_cls=config["loader_cls"],
                loader_kwargs=config["loader_kwargs"],
                show_progress=True,
            )
            docs = loader.load()
            if docs:
                # 增强元数据
                if enricher:
                    docs = enricher.enrich_documents(docs, docs[0].metadata.get("source", ""))
                all_docs.extend(docs)
                print(f"  ✅ {ext}: 加载 {len(docs)} 个文档")
        except Exception as e:
            print(f"  ⚠️ {ext} 加载失败: {e}")

    # 2. 加载 .docx 文件
    docx_files = []
    for root, dirs, files in os.walk(data_dir):
        for file in files:
            if file.endswith(".docx"):
                docx_files.append(os.path.join(root, file))

    for docx_path in docx_files:
        try:
            docs = load_docx(docx_path)
            if enricher:
                docs = enricher.enrich_documents(docs, docx_path)
            all_docs.extend(docs)
            print(f"  ✅ .docx: 加载 {docx_path}")
        except Exception as e:
            print(f"  ⚠️ .docx 加载失败 {docx_path}: {e}")

    # 3. 加载 .xlsx 文件
    xlsx_files = []
    for root, dirs, files in os.walk(data_dir):
        for file in files:
            if file.endswith(".xlsx"):
                xlsx_files.append(os.path.join(root, file))

    for xlsx_path in xlsx_files:
        try:
            docs = load_xlsx(xlsx_path)
            if enricher:
                docs = enricher.enrich_documents(docs, xlsx_path)
            all_docs.extend(docs)
            print(f"  ✅ .xlsx: 加载 {xlsx_path} ({len(docs)}个工作表)")
        except Exception as e:
            print(f"  ⚠️ .xlsx 加载失败 {xlsx_path}: {e}")

    return all_docs


def split_documents(docs: List[Document], chunk_size: int = 800, 
                    chunk_overlap: int = 100, enrich_metadata: bool = True) -> List[Document]:
    """
    将文档切分为文本块
    
    Args:
        docs: Document 列表
        chunk_size: 块大小
        chunk_overlap: 块重叠
        enrich_metadata: 是否增强元数据
    
    Returns:
        切分后的 Document 列表
    """
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n\n", "\n\n", "\n", "。", ".", " ", ""],
    )
    chunks = text_splitter.split_documents(docs)

    # 为每个 chunk 添加来源文件名和增强元数据
    enricher = MetadataEnricher() if enrich_metadata else None
    
    for chunk in chunks:
        source = chunk.metadata.get("source", "unknown")
        chunk.metadata["source_file"] = os.path.basename(source)
        
        # 如果需要增强元数据且没有已增强的元数据
        if enricher and "doc_type" not in chunk.metadata:
            chunk = enricher.enrich_metadata(chunk, source)

    return chunks


def load_and_split(file_path: str, chunk_size: int = 800, 
                   chunk_overlap: int = 100, enrich_metadata: bool = True) -> List[Document]:
    """
    加载单个文件并切分
    
    Args:
        file_path: 文件路径
        chunk_size: 块大小
        chunk_overlap: 块重叠
        enrich_metadata: 是否增强元数据
    
    Returns:
        切分后的 Document 列表
    """
    docs = load_single_document(file_path)
    if not docs:
        return []
    
    return split_documents(docs, chunk_size, chunk_overlap, enrich_metadata)


def main():
    data_dir = "data"

    print("=" * 50)
    print("🚀 文档摄入管道 v2 (元数据增强)")
    print("=" * 50)

    # 1. 加载文档
    docs = load_documents(data_dir, enrich_metadata=True)
    if not docs:
        print("⚠️ 未找到任何文档。")
        return

    print(f"\n✅ 共加载 {len(docs)} 个文档")

    # 2. 切片
    print("✂️  正在切片...")
    chunks = split_documents(docs, enrich_metadata=False)  # 已经在 load_documents 中增强过
    print(f"✅ 共生成 {len(chunks)} 个文本块")

    # 3. 预览
    if chunks:
        print("\n--- 首个文本块预览 ---")
        print(f"📄 来源: {chunks[0].metadata.get('source_file', 'unknown')}")
        print(f"📝 类型: {chunks[0].metadata.get('doc_type', 'unknown')}")
        print(f"🏗️  项目: {chunks[0].metadata.get('project', 'unknown')}")
        print(f"📝 内容: {chunks[0].page_content[:200]}...")
        print("----------------------\n")

    print(f"🎉 文档摄入完成！共 {len(chunks)} 个文本块待向量化")
    print(f"💡 下一步: 运行 python vectorstore_setup.py 进行向量化")


if __name__ == "__main__":
    main()
